from pathlib import Path
import pickle

import faiss
from pypdf import PdfReader
from docx import Document
from sentence_transformers import SentenceTransformer


# ============================================================
# CONFIGURATION
# ============================================================

KNOWLEDGE_DIR = Path("knowledge")
INDEX_DIR = Path("rag_index")

MODEL_NAME = "all-MiniLM-L6-v2"

CHUNK_SIZE = 1000
CHUNK_OVERLAP = 150


# ============================================================
# DOCUMENT LOADING
# ============================================================

def load_documents():
    documents = []

    if not KNOWLEDGE_DIR.exists():
        raise RuntimeError(
            f"Knowledge directory does not exist: {KNOWLEDGE_DIR}"
        )

    files = list(KNOWLEDGE_DIR.rglob("*"))

    for path in files:

        if not path.is_file():
            continue

        suffix = path.suffix.lower()

        try:

            # ==================================================
            # PDF
            # ==================================================

            if suffix == ".pdf":

                reader = PdfReader(str(path))

                for page_number, page in enumerate(reader.pages):

                    text = page.extract_text() or ""
                    text = text.strip()

                    if text:

                        documents.append(
                            {
                                "text": text,
                                "source": path.name,
                                "page": page_number + 1,
                            }
                        )

            # ==================================================
            # DOCX
            # ==================================================

            elif suffix == ".docx":

                document = Document(str(path))

                parts = []

                # ----------------------------------------------
                # Paragraphs
                # ----------------------------------------------

                for paragraph in document.paragraphs:

                    text = paragraph.text.strip()

                    if text:
                        parts.append(text)

                # ----------------------------------------------
                # Tables
                # ----------------------------------------------

                for table in document.tables:

                    for row in table.rows:

                        cells = []

                        for cell in row.cells:

                            text = cell.text.strip()

                            if text:
                                cells.append(text)

                        if cells:
                            parts.append(" | ".join(cells))

                text = "\n".join(parts).strip()

                if text:

                    documents.append(
                        {
                            "text": text,
                            "source": path.name,
                            "page": None,
                        }
                    )

            # ==================================================
            # TXT / MARKDOWN
            # ==================================================

            elif suffix in {".txt", ".md"}:

                text = path.read_text(
                    encoding="utf-8",
                    errors="ignore",
                ).strip()

                if text:

                    documents.append(
                        {
                            "text": text,
                            "source": path.name,
                            "page": None,
                        }
                    )

        except Exception as exc:

            print(
                f"WARNING: could not read "
                f"{path}: {exc}"
            )

    return documents


# ============================================================
# CHUNKING
# ============================================================

def chunk_text(text: str):

    # Normalize whitespace
    text = " ".join(text.split())

    chunks = []

    start = 0

    while start < len(text):

        end = start + CHUNK_SIZE

        chunk = text[start:end].strip()

        if chunk:
            chunks.append(chunk)

        if end >= len(text):
            break

        start = end - CHUNK_OVERLAP

    return chunks


# ============================================================
# BUILD CHUNKS
# ============================================================

def build_chunks(documents):

    chunks = []

    for document in documents:

        document_chunks = chunk_text(
            document["text"]
        )

        for chunk in document_chunks:

            chunks.append(
                {
                    "text": chunk,
                    "source": document["source"],
                    "page": document["page"],
                }
            )

    return chunks


# ============================================================
# BUILD FAISS INDEX
# ============================================================

def build_index(chunks):

    print()
    print(
        f"Loading embedding model: {MODEL_NAME}"
    )

    model = SentenceTransformer(
        MODEL_NAME
    )

    print()
    print("Generating embeddings...")

    texts = [
        chunk["text"]
        for chunk in chunks
    ]

    embeddings = model.encode(
        texts,
        show_progress_bar=True,
        normalize_embeddings=True,
    )

    dimension = embeddings.shape[1]

    print()
    print(
        f"Embedding dimension: {dimension}"
    )

    # --------------------------------------------------------
    # FAISS
    # --------------------------------------------------------

    index = faiss.IndexFlatIP(
        dimension
    )

    index.add(embeddings)

    return index


# ============================================================
# SAVE INDEX
# ============================================================

def save_index(index, chunks):

    INDEX_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    index_path = (
        INDEX_DIR / "index.faiss"
    )

    chunks_path = (
        INDEX_DIR / "chunks.pkl"
    )

    faiss.write_index(
        index,
        str(index_path),
    )

    with open(
        chunks_path,
        "wb",
    ) as file:

        pickle.dump(
            chunks,
            file,
        )

    print()
    print(
        f"FAISS index saved to: {index_path}"
    )

    print(
        f"Chunk metadata saved to: {chunks_path}"
    )


# ============================================================
# MAIN
# ============================================================

def main():

    print()
    print(
        "========================================"
    )
    print(
        " Hyperion RAG index builder"
    )
    print(
        "========================================"
    )
    print()

    # --------------------------------------------------------
    # Load documents
    # --------------------------------------------------------

    print(
        f"Reading documents from: {KNOWLEDGE_DIR}"
    )

    documents = load_documents()

    if not documents:

        raise RuntimeError(
            "No supported documents were found "
            "in knowledge/."
        )

    print()
    print(
        f"Loaded {len(documents)} document sections."
    )

    # --------------------------------------------------------
    # Show documents
    # --------------------------------------------------------

    print()
    print("Documents loaded:")

    unique_sources = sorted(
        {
            document["source"]
            for document in documents
        }
    )

    for source in unique_sources:

        print(
            f"  - {source}"
        )

    # --------------------------------------------------------
    # Create chunks
    # --------------------------------------------------------

    print()
    print("Creating chunks...")

    chunks = build_chunks(
        documents
    )

    if not chunks:

        raise RuntimeError(
            "No text chunks were generated "
            "from the documents."
        )

    print(
        f"Created {len(chunks)} chunks."
    )

    # --------------------------------------------------------
    # Build vector index
    # --------------------------------------------------------

    index = build_index(
        chunks
    )

    # --------------------------------------------------------
    # Save
    # --------------------------------------------------------

    save_index(
        index,
        chunks,
    )

    # --------------------------------------------------------
    # Summary
    # --------------------------------------------------------

    print()
    print(
        "========================================"
    )
    print(
        " RAG index created successfully"
    )
    print(
        "========================================"
    )
    print()

    print(
        f"Documents: {len(unique_sources)}"
    )

    print(
        f"Sections:  {len(documents)}"
    )

    print(
        f"Chunks:    {len(chunks)}"
    )

    print()
    print(
        f"Index directory: {INDEX_DIR}"
    )

    print()


# ============================================================
# ENTRY POINT
# ============================================================

if __name__ == "__main__":
    main()